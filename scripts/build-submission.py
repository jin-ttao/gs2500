#!/usr/bin/env python3
"""Build the five-page GS2500 deck, PDF, images, and captioned demo video."""

from __future__ import annotations

import json
import math
import shutil
import subprocess
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageEnhance, ImageFont
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
SLIDES = OUT / "slides"
MEDIA = OUT / ".video-parts"
FONT = "/System/Library/Fonts/AppleSDGothicNeo.ttc"
WIDTH, HEIGHT = 1920, 1080

INK = "#10233F"
MUTED = "#637083"
PAPER = "#F5F7FA"
WHITE = "#FFFFFF"
BLUE = "#1C53D8"
SKY = "#DCE8FF"
GREEN = "#04A866"
MINT = "#D8F4E8"
ORANGE = "#FF7A3D"
AMBER = "#FFEEE6"
RED = "#D84B4B"
PURPLE = "#805AD5"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT, size, index=1 if bold else 0)


def rr(draw: ImageDraw.ImageDraw, box, radius=24, fill=WHITE, outline=None, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def text(draw, xy, value, size, color=INK, bold=False, anchor=None, spacing=8):
    draw.multiline_text(xy, value, font=font(size, bold), fill=color, anchor=anchor, spacing=spacing)


def fit_cover(image: Image.Image, box) -> Image.Image:
    x1, y1, x2, y2 = box
    width, height = x2 - x1, y2 - y1
    ratio = max(width / image.width, height / image.height)
    resized = image.resize((round(image.width * ratio), round(image.height * ratio)), Image.Resampling.LANCZOS)
    left = max(0, (resized.width - width) // 2)
    top = max(0, (resized.height - height) // 2)
    return resized.crop((left, top, left + width, top + height))


def base(slide_no: int, label: str, title: str, subtitle: str = ""):
    image = Image.new("RGB", (WIDTH, HEIGHT), PAPER)
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, WIDTH, 18), fill=GREEN)
    text(draw, (86, 70), f"GS2500  /  {label}", 26, BLUE, True)
    text(draw, (1830, 72), f"0{slide_no} / 05", 24, MUTED, True, "ra")
    text(draw, (86, 132), title, 58, INK, True)
    if subtitle:
        text(draw, (88, 214), subtitle, 26, MUTED)
    return image, draw


def paste_rounded(canvas_image, source, box, radius=28):
    fitted = fit_cover(source.convert("RGB"), box)
    mask = Image.new("L", fitted.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, *fitted.size), radius=radius, fill=255)
    canvas_image.paste(fitted, (box[0], box[1]), mask)


def money(value: int) -> str:
    return f"{value / 1_000_000:.3f}백만원"


def slide_one(current_photo: Image.Image):
    image = Image.new("RGB", (WIDTH, HEIGHT), INK)
    photo = fit_cover(current_photo, (1020, 0, WIDTH, HEIGHT))
    photo = ImageEnhance.Brightness(photo).enhance(0.78)
    image.paste(photo, (1020, 0))
    overlay = Image.new("RGBA", (WIDTH, HEIGHT), (0, 0, 0, 0))
    od = ImageDraw.Draw(overlay)
    for i in range(520):
        alpha = int(255 * (1 - i / 520))
        od.line((760 + i, 0, 760 + i, HEIGHT), fill=(16, 35, 63, alpha))
    image = Image.alpha_composite(image.convert("RGBA"), overlay).convert("RGB")
    draw = ImageDraw.Draw(image)
    rr(draw, (82, 76, 326, 126), 24, GREEN)
    text(draw, (204, 101), "GS2500", 25, WHITE, True, "mm")
    text(draw, (82, 200), "점포마다 다른 답을\n작게 검증합니다", 72, WHITE, True, spacing=18)
    text(draw, (84, 402), "매니저의 판단부터 점주의 작은 실행,\n4주 뒤 다음 지원까지 한 화면 흐름으로 연결", 30, "#D9E2EF", spacing=12)
    stages = ["전년 동기비", "24시간 실험", "매니저 승인", "점주 실행", "4주 회고"]
    x = 84
    for index, stage in enumerate(stages):
        rr(draw, (x, 570, x + 160, 648), 18, "#193252", "#315176", 2)
        text(draw, (x + 80, 609), stage, 20, WHITE, True, "mm")
        x += 180
        if index < len(stages) - 1:
            text(draw, (x - 10, 609), "›", 31, GREEN, True, "mm")
    rr(draw, (84, 758, 884, 920), 26, "#17304F", "#345574", 2)
    text(draw, (116, 790), "발표 경로", 22, GREEN, True)
    text(draw, (116, 835), "API 키 없이 공개 URL에서 끝까지 동작", 31, WHITE, True)
    text(draw, (116, 880), "합성 데이터 · 실제 전송/발주/인과 주장 없음", 23, "#C7D4E3")
    text(draw, (84, 1008), "GS그룹 x 랄프톤 해커톤  ·  2026.09.22", 22, "#9DB0C7")
    return image


def slide_two(stores):
    image, draw = base(2, "MANAGER", "어디에 개입할지 먼저 고릅니다", "전년 동기비는 작성된 합성 실적이며, 24시간 실험과 다른 지표입니다.")
    rr(draw, (84, 290, 1260, 980), 30, WHITE)
    text(draw, (120, 326), "담당 9개 점포 · 전년 동기비", 27, INK, True)
    max_abs = 9.0
    zero_x = 700
    draw.line((zero_x, 390, zero_x, 916), fill="#B8C2D0", width=3)
    for i, store in enumerate(stores):
        y = 408 + i * 57
        yoy = store["yoy"]
        text(draw, (130, y), store["name"], 22, INK, True)
        text(draw, (430, y), store["format"], 17, MUTED)
        length = abs(yoy) / max_abs * 420
        if yoy < 0:
            draw.rounded_rectangle((zero_x - length, y + 3, zero_x, y + 29), radius=13, fill=ORANGE if store["id"] == "samsung" else "#EFA182")
        else:
            draw.rounded_rectangle((zero_x, y + 3, zero_x + length, y + 29), radius=13, fill=GREEN)
        text(draw, (1140, y + 16), f"{yoy:+.1f}%", 22, RED if yoy < 0 else GREEN, True, "rm")
    rr(draw, (1305, 290, 1838, 598), 30, INK)
    text(draw, (1344, 330), "우선 확인", 22, "#9EC4FF", True)
    text(draw, (1344, 378), "삼성역점", 42, WHITE, True)
    text(draw, (1344, 448), "전년 동기비", 20, "#ACB8C8")
    text(draw, (1344, 486), "-8.4%", 56, ORANGE, True)
    text(draw, (1580, 486), "42.84백만원", 26, WHITE, True)
    text(draw, (1344, 558), "다음: B-03 행사 매대 확인", 22, "#D8E2EE")
    rr(draw, (1305, 630, 1838, 980), 30, WHITE)
    text(draw, (1344, 670), "진열 외 가설도 남깁니다", 27, INK, True)
    items = [("발주", "저녁 결품이 먼저인 점포"), ("행사", "가격/프로모션 조건이 다른 점포"), ("현장 지원", "인력과 매대 여유가 부족한 점포")]
    for i, (name, note) in enumerate(items):
        y = 740 + i * 74
        rr(draw, (1342, y, 1444, y + 44), 15, MINT if i == 2 else SKY)
        text(draw, (1393, y + 22), name, 18, GREEN if i == 2 else BLUE, True, "mm")
        text(draw, (1465, y + 22), note, 18, MUTED, False, "lm")
    return image


def slide_three(candidates):
    image, draw = base(3, "24H 3D EXPERIMENT", "같은 조건의 네 후보를 실제 계산했습니다", "동일 시드 11 · 초기 총재고 동일 · 잠재 고객 1,000명 일정 동일 · 로컬 규칙 엔진")
    rr(draw, (84, 292, 1265, 940), 30, WHITE)
    text(draw, (122, 332), "삼성역점 A/B/C/D · 24시간 모의 결제매출", 27, INK, True)
    colors = {"A": BLUE, "B": GREEN, "C": ORANGE, "D": PURPLE}
    titles = {c["candidateId"]: c["title"] for c in candidates}
    max_rev = max(c["paidRevenue"] for c in candidates)
    for i, c in enumerate(candidates):
        cid = c["candidateId"]
        y = 430 + i * 112
        rr(draw, (130, y - 6, 195, y + 59), 18, colors[cid])
        text(draw, (162, y + 26), cid, 29, WHITE, True, "mm")
        text(draw, (222, y), titles[cid], 24, INK, True)
        bar_width = round(c["paidRevenue"] / max_rev * 700)
        draw.rounded_rectangle((472, y + 2, 472 + bar_width, y + 43), radius=20, fill=colors[cid])
        text(draw, (1208, y + 22), money(c["paidRevenue"]), 23, INK, True, "rm")
        text(draw, (222, y + 42), f"구매자 {c['day']['buyers']}명 · 결품 수요 {c['day']['totalStockoutDemand']}건", 17, MUTED)
    rr(draw, (1305, 292, 1838, 602), 30, INK)
    best = max(candidates, key=lambda c: c["paidRevenue"])
    text(draw, (1344, 332), "절대 매출 최상위", 22, "#9FE4C7", True)
    text(draw, (1344, 382), f"후보 {best['candidateId']} · {best['title']}", 33, WHITE, True)
    text(draw, (1344, 454), money(best["paidRevenue"]), 54, GREEN, True)
    text(draw, (1344, 526), f"매출총이익 {best['paidGrossProfit']:,}원", 23, "#D5E0EC")
    text(draw, (1344, 562), "실제 기록 24초 재생", 21, "#9EC4FF")
    rr(draw, (1305, 638, 1838, 940), 30, AMBER)
    text(draw, (1344, 676), "해석 경계", 23, ORANGE, True)
    rules = ["A는 본사 표준 후보", "A는 현재 진열이 아님", "현재 대비 상승률을 주장하지 않음", "완료 기록을 사전 계산해 재생"]
    for i, rule in enumerate(rules):
        text(draw, (1348, 734 + i * 48), "●", 15, ORANGE, True)
        text(draw, (1382, 728 + i * 48), rule, 22, INK, i == 1)
    text(draw, (84, 994), "점포 간 절대 매출은 상권, 재고 배율, 시드가 달라 직접 순위로 쓰지 않습니다.", 22, MUTED)
    return image


def slide_four(current_photo: Image.Image, after_photo: Image.Image):
    image, draw = base(4, "MANAGER → STORE OWNER", "승인된 작은 변경만 점주에게 보입니다", "후보 B 전체 결과를 참고해 2개 상품 이동만 제안합니다. 이 부분 효과는 별도 계산값이 아닙니다.")
    paste_rounded(image, current_photo, (84, 310, 724, 744), 28)
    paste_rounded(image, after_photo, (764, 310, 1404, 744), 28)
    text(draw, (110, 334), "합성 현재 이미지", 21, WHITE, True)
    text(draw, (790, 334), "발표용 합성 변경 사진", 21, WHITE, True)
    rr(draw, (140, 650, 430, 720), 19, "#10233FDD")
    text(draw, (285, 685), "1단 → 3단 · 제로 음료", 21, WHITE, True, "mm")
    rr(draw, (838, 650, 1176, 720), 19, "#10233FDD")
    text(draw, (1007, 685), "3단 → 2단 · 담백 크래커", 21, WHITE, True, "mm")
    rr(draw, (1440, 310, 1838, 744), 30, WHITE)
    text(draw, (1478, 350), "점주가 받는 정보", 27, INK, True)
    cards = [("12분", "예상 작업"), ("발주 없음", "신규 발주 아님"), ("수정 가능", "안내 문자 초안")]
    for i, (value, label) in enumerate(cards):
        y = 420 + i * 88
        rr(draw, (1476, y, 1800, y + 70), 18, MINT if i == 0 else SKY)
        text(draw, (1500, y + 18), value, 24, GREEN if i == 0 else BLUE, True)
        text(draw, (1774, y + 36), label, 17, MUTED, False, "rm")
    text(draw, (1478, 692), "수용 · 부분 실행 · 거절", 22, INK, True)
    rr(draw, (84, 784, 1838, 960), 30, INK)
    choices = [("수용", "제안대로"), ("부분 실행", "가능한 만큼"), ("거절", "이유만 남김"), ("사진", "선택 사항")]
    x = 122
    for label, note in choices:
        rr(draw, (x, 824, x + 340, 916), 22, "#183352", "#345776", 2)
        text(draw, (x + 26, 844), label, 25, GREEN, True)
        text(draw, (x + 26, 882), note, 20, WHITE)
        x += 380
    text(draw, (84, 1005), "사진은 로컬 미리보기와 메타데이터 접수만 제공합니다. 자동 판독, 진열 검증, 서버 저장은 없습니다.", 21, MUTED)
    return image


def slide_five():
    image, draw = base(5, "4 WEEKS LATER", "좋은 사례, 부진, 미회신을 같은 분모로 봅니다", "별도 4주 합성 관찰입니다. 24시간 실험값도, 인과적으로 입증된 매출 상승도 아닙니다.")
    labels = [("제안", "4", "검토 대상"), ("승인", "3", "제안 4"), ("회신", "2", "승인 3"), ("사진", "2", "회신 2")]
    for i, (label, value, denom) in enumerate(labels):
        x = 84 + i * 315
        rr(draw, (x, 300, x + 280, 492), 28, WHITE)
        text(draw, (x + 28, 332), label, 23, MUTED, True)
        text(draw, (x + 28, 376), value, 66, BLUE if i < 2 else GREEN, True)
        text(draw, (x + 28, 454), f"분모: {denom}", 18, MUTED)
    rr(draw, (1380, 300, 1838, 492), 28, INK)
    text(draw, (1416, 332), "전국 확대 결론", 22, "#9EC4FF", True)
    text(draw, (1416, 380), "아직 보류", 48, WHITE, True)
    text(draw, (1416, 452), "네 사례로 확정하지 않음", 19, "#C8D4E2")
    cases = [
        ("좋은 관찰", "삼성역점", "다른 점포에서 재검증", GREEN, MINT),
        ("부진", "역세권 데모점", "발주 주기 우선 검토", ORANGE, AMBER),
        ("미회신", "주거지 데모점", "현장 인력과 여유 확인", BLUE, SKY),
    ]
    for i, (status, store, next_step, color, pale) in enumerate(cases):
        x = 84 + i * 584
        rr(draw, (x, 536, x + 548, 842), 28, WHITE)
        rr(draw, (x + 28, 568, x + 178, 614), 16, pale)
        text(draw, (x + 103, 591), status, 19, color, True, "mm")
        text(draw, (x + 28, 648), store, 29, INK, True)
        text(draw, (x + 28, 704), "다른 설명 가능성", 18, MUTED, True)
        alt = ["행사 · 날씨 · 자연 변동", "결품 · 퇴근 시간 수요", "주말 인력 부족 · 실행 불명"][i]
        text(draw, (x + 28, 740), alt, 20, INK)
        rr(draw, (x + 28, 786, x + 518, 824), 14, INK)
        text(draw, (x + 273, 805), next_step, 18, WHITE, True, "mm")
    rr(draw, (84, 878, 1838, 998), 26, "#E7ECF3")
    text(draw, (116, 910), "다음 지원", 21, BLUE, True)
    text(draw, (300, 910), "재검증 후보 선정", 22, INK, True)
    text(draw, (690, 910), "부분 실행과 미회신 점포의 현장 지원", 22, INK, True)
    text(draw, (1340, 910), "결품 점포는 발주 가설로 전환", 22, INK, True)
    text(draw, (116, 966), "공개 앱 · 발표 자료 · 영상 · 원문 · 세션 기록: gs2500.vercel.app/deliverables/", 20, MUTED)
    return image


NARRATION = [
    "GS2500은 점포마다 다른 진열 답을 작은 실험과 실행으로 연결합니다. 영업관리 매니저가 전년 동기비를 보고 개입할 점포를 고르고, 같은 조건의 스물네 시간 삼차원 실험을 비교합니다. 승인한 작은 변경만 점주에게 보이며, 점주는 가능한 만큼 실행한 뒤 선택과 사진을 남깁니다. 네 주 뒤에는 좋은 사례만 골라 보지 않고 부진과 미회신을 함께 봅니다. 이 데모는 합성 데이터이고 실제 문자, 발주, 계정, 사내 시스템과 연결되지 않습니다.",
    "매니저 홈은 담당 아홉 점포의 전년 동기비부터 보여 줍니다. 삼성역점은 마이너스 팔 점 사 퍼센트로 우선 확인 대상입니다. 이 수치는 별도로 작성한 합성 실적이며 뒤의 스물네 시간 실험과 같은 숫자가 아닙니다. 매니저는 매대 회전 둔화가 있는 비 영 삼 행사 매대를 고릅니다. 진열이 맞지 않는 점포에는 발주, 행사, 현장 지원 같은 다른 가설을 남깁니다. 진열은 여러 개입 수단 중 하나입니다.",
    "삼성역점의 네 후보는 시드 십일, 같은 초기 총재고, 같은 합성 페르소나 천 명의 잠재 고객 일정을 공유합니다. 로컬 규칙 엔진을 스물네 시간 끝까지 실행한 절대 결과는 에이 백육십사 점 육삼만 원, 비 백팔십오 점 칠삼만 원, 씨 백육십이 점 일오만 원, 디 백오십이 점 사사만 원입니다. 비가 이 실행에서는 가장 높습니다. 에이는 본사 표준 후보일 뿐 현재 진열이 아닙니다. 현재 대비 상승률은 주장하지 않습니다. 발표에서는 완료 기록을 이십사 초로 재생합니다.",
    "매니저는 비 후보 전체를 그대로 점주에게 떠넘기지 않습니다. 현재 합성 이미지와 재고, 작성된 합성 실적 근거를 보고 제로 음료를 일 단에서 삼 단으로, 담백 크래커를 삼 단에서 이 단으로 옮기는 작은 안을 승인합니다. 약 십이 분, 새 발주 없음, 되돌릴 수 있는 변경입니다. 이 두 상품만 옮긴 효과를 따로 계산한 값은 아닙니다. 점주는 수용, 부분 실행, 거절을 고르고 선택적 의견과 사진을 남깁니다. 사진은 로컬 미리보기이며 자동 판독하거나 서버에 저장하지 않습니다.",
    "네 주 뒤 회고는 제안 네 건, 승인 세 건, 회신 두 건, 사진 두 건처럼 분모를 나눠 봅니다. 좋은 관찰은 다른 점포에서 다시 검증하고, 부진 사례는 진열보다 발주 주기와 결품이 큰 설명인지 확인합니다. 미회신 점포는 성과에서 빼고 현장 인력과 실행 여부를 먼저 확인합니다. 이 화면은 별도 합성 관찰이며 인과적으로 입증된 매출 상승이 아닙니다. 공개 주소에서 전체 경로와 기존 삼차원 실험실, 발표 자료, 영상, 실제 목표 원문과 세션 기록을 확인할 수 있습니다.",
]


def save_slides(slide_images):
    SLIDES.mkdir(parents=True, exist_ok=True)
    for path in SLIDES.glob("*.png"):
        path.unlink()
    paths = []
    for i, image in enumerate(slide_images, 1):
        path = SLIDES / f"{i:02d}.png"
        image.save(path, optimize=True)
        paths.append(path)
    return paths


def build_pptx(paths):
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    for path in paths:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(str(path), 0, 0, width=presentation.slide_width, height=presentation.slide_height)
    presentation.save(OUT / "gs2500-presentation.pptx")


def build_pdf(paths):
    output = OUT / "gs2500-presentation.pdf"
    pdf = canvas.Canvas(str(output), pagesize=(WIDTH, HEIGHT))
    for path in paths:
        pdf.drawImage(str(path), 0, 0, WIDTH, HEIGHT)
        pdf.showPage()
    pdf.save()


def run(command):
    subprocess.run(command, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def build_video(paths):
    if not shutil.which("ffmpeg"):
        print("ffmpeg is unavailable; deck was still created.", file=sys.stderr)
        return None
    if MEDIA.exists():
        shutil.rmtree(MEDIA)
    MEDIA.mkdir(parents=True)
    segments = []
    for i, path in enumerate(paths, 1):
        segment = MEDIA / f"{i:02d}.mp4"
        run([
            "ffmpeg", "-y", "-loop", "1", "-framerate", "1", "-i", str(path),
            "-f", "lavfi", "-i", "anullsrc=r=48000:cl=stereo", "-t", "18",
            "-vf", "scale=1280:720,format=yuv420p", "-c:v", "libx264", "-preset", "ultrafast", "-tune", "stillimage",
            "-threads", "2", "-c:a", "aac", "-b:a", "96k", "-shortest", "-movflags", "+faststart", str(segment),
        ])
        segments.append(segment)
    concat = MEDIA / "concat.txt"
    concat.write_text("".join(f"file '{p.as_posix()}'\n" for p in segments), encoding="utf-8")
    output = OUT / "gs2500-demo-video.mp4"
    run(["ffmpeg", "-y", "-f", "concat", "-safe", "0", "-i", str(concat), "-c", "copy", "-movflags", "+faststart", str(output)])
    shutil.rmtree(MEDIA)
    return output


def build_index(video_duration: float | None):
    duration = "확인 중" if video_duration is None else f"{math.floor(video_duration / 60)}분 {round(video_duration % 60)}초"
    html = f"""<!doctype html>
<html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>GS2500 제출물</title><style>
*{{box-sizing:border-box}}body{{margin:0;background:#f5f7fa;color:#10233f;font-family:-apple-system,BlinkMacSystemFont,'Apple SD Gothic Neo',sans-serif}}main{{max-width:1180px;margin:0 auto;padding:64px 32px 100px}}a{{color:#1c53d8}}.eyebrow{{color:#1c53d8;font-weight:800;letter-spacing:.08em}}h1{{font-size:52px;margin:12px 0}}.lede{{font-size:21px;color:#637083;line-height:1.7;max-width:820px}}.links{{display:grid;grid-template-columns:repeat(2,1fr);gap:16px;margin:36px 0}}.card{{display:block;padding:24px;border-radius:20px;background:#fff;text-decoration:none;color:#10233f;box-shadow:0 8px 30px #10233f10}}.card b{{display:block;font-size:22px;margin-bottom:8px}}.card span{{color:#637083}}video{{width:100%;border-radius:24px;background:#10233f;box-shadow:0 20px 50px #10233f22}}h2{{margin-top:56px}}.slides{{display:grid;gap:20px}}.slides img{{width:100%;border-radius:18px;box-shadow:0 10px 30px #10233f18}}.truth{{padding:24px 28px;border-radius:18px;background:#e8f6ef;line-height:1.7}}@media(max-width:800px){{.links{{grid-template-columns:1fr}}h1{{font-size:38px}}}}
</style></head><body><main><span class="eyebrow">GS2500 · HACKATHON SUBMISSION</span><h1>공개 제출물</h1>
<p class="lede">점포 성과 확인부터 같은 조건의 24시간 실험, 매니저 승인, 점주의 작은 실행, 4주 합성 회고까지 이어지는 데스크톱 데모입니다.</p>
<div class="links"><a class="card" href="../"><b>데스크톱 웹앱 열기</b><span>3분 발표 경로 시작</span></a><a class="card" href="./gs2500-presentation.pdf"><b>5페이지 발표 자료</b><span>PDF 다운로드</span></a><a class="card" href="./gs2500-presentation.pptx"><b>발표 자료 원본</b><span>PPTX 다운로드</span></a><a class="card" href="./goal-original.md"><b>실제 /goal 원문</b><span>입력한 목표 전문</span></a><a class="card" href="./session/main-session.jsonl"><b>메인 세션 JSONL</b><span>Codex 허브 세션 기록</span></a><a class="card" href="https://github.com/jin-ttao/gs2500"><b>공개 GitHub</b><span>소스와 재현 절차</span></a></div>
<h2>무음 자막형 데모 영상 · {duration}</h2><video controls preload="metadata" poster="./slides/01.png"><source src="./gs2500-demo-video.mp4" type="video/mp4"></video><p class="lede">슬라이드 안의 자막만으로 볼 수 있으며, 발표자가 읽을 한국어 스크립트는 <a href="./narration.txt">narration.txt</a>에 있습니다.</p>
<p class="truth"><b>표시 구분</b><br>전년 동기비는 작성된 합성 실적, A/B/C/D는 24시간 로컬 규칙 엔진 계산, 3D 행동은 완료 기록의 사전 재생, 회고는 별도 4주 합성 관찰입니다. 실제 GS25 데이터, 고객 행동 예측, 인과적으로 입증된 매출 상승을 뜻하지 않습니다.</p>
<h2>발표 자료 미리보기</h2><div class="slides">{''.join(f'<img src="./slides/{i:02d}.png" alt="GS2500 발표 자료 {i}쪽">' for i in range(1, 6))}</div>
</main></body></html>"""
    (OUT / "index.html").write_text(html, encoding="utf-8")


def ffprobe_duration(path: Path | None):
    if path is None or not path.exists() or not shutil.which("ffprobe"):
        return None
    value = subprocess.check_output(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "default=nw=1:nk=1", str(path)], text=True)
    return float(value.strip())


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    data = json.loads((ROOT / "app/data/experiment-results.json").read_text(encoding="utf-8"))
    samsung = next(store for store in data["stores"] if store["id"] == "samsung")
    stores = [
        {"id": "samsung", "name": "삼성역점", "format": "합성 오피스형", "yoy": -8.4},
        {"id": "station", "name": "역세권 데모점", "format": "합성 세로형", "yoy": -6.2},
        {"id": "residential", "name": "주거지 데모점", "format": "합성 주거지형", "yoy": -4.9},
        {"id": "compact", "name": "골목 데모점", "format": "합성 소형점", "yoy": -3.3},
        {"id": "riverside", "name": "한강변 데모점", "format": "합성 수변형", "yoy": -1.7},
        {"id": "park", "name": "공원 데모점", "format": "합성 산책형", "yoy": 0.4},
        {"id": "university", "name": "대학가 데모점", "format": "합성 캠퍼스형", "yoy": 1.1},
        {"id": "tourism", "name": "관광지 데모점", "format": "합성 관광형", "yoy": 2.8},
        {"id": "cafe", "name": "카페 데모점", "format": "합성 취식형", "yoy": 4.6},
    ]
    current = Image.open(ROOT / "app/assets/samsung-b03-current-synthetic.png")
    after = Image.open(ROOT / "app/assets/samsung-b03-after-synthetic.png")
    paths = save_slides([
        slide_one(current), slide_two(stores), slide_three(samsung["candidates"]),
        slide_four(current, after), slide_five(),
    ])
    build_pptx(paths)
    build_pdf(paths)
    (OUT / "narration.txt").write_text("\n\n".join(f"{i}. {item}" for i, item in enumerate(NARRATION, 1)) + "\n", encoding="utf-8")
    video = build_video(paths)
    duration = ffprobe_duration(video)
    if duration and duration > 180:
        raise RuntimeError(f"Demo video exceeds three minutes: {duration:.2f}s")
    build_index(duration)
    print(f"Built 5 slides, PDF, PPTX, and video ({duration:.2f}s).")


if __name__ == "__main__":
    main()
